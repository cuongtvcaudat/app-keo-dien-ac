#!/usr/bin/env python3
"""Tao file Excel mau (du_lieu.xlsx) voi dung cot + vai dong vi du."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(HERE, "data", "du_lieu.xlsx")

# Cot: 'file' = ten anh trong data/input (khong can duoi .jpg). 'template' = kieu_A/kieu_B.
HEADERS = ["file", "template", "ngay", "lat", "lon", "bearing_deg", "huong",
           "dia_chi", "phuong", "huyen", "tinh", "altitude", "speed",
           "ma_tram", "ghi_chu", "index"]

# Du lieu lay tu 5 anh mau (de test / lam vi du).
SAMPLES = [
    {"file": "LDG0911", "template": "kieu_A", "ngay": "22 Th7, 2026 16:43:31",
     "lat": "11.51856895856612N", "lon": "107.83645898044976E",
     "phuong": "B'Lao", "tinh": "Lâm Đồng", "ma_tram": "LDG0911"},
    {"file": "LDG0772", "template": "kieu_B", "ngay": "14/07/2026",
     "lat": "11.71146542", "lon": "108.37751334", "bearing_deg": "34", "huong": "34° NE",
     "dia_chi": "19 Nguyễn Bá Ngọc", "huyen": "Đức Trọng", "tinh": "Lâm Đồng",
     "altitude": "913.2", "speed": "0.0", "ma_tram": "LDG0772 M1", "ghi_chu": ".t3", "index": "49"},
    {"file": "LDG0747_m3_t3", "template": "kieu_B", "ngay": "15/07/2026",
     "lat": "11.69125766", "lon": "108.3746989", "bearing_deg": "258", "huong": "258° W",
     "huyen": "Đức Trọng", "tinh": "Lâm Đồng", "altitude": "903.0", "speed": "0.0",
     "ma_tram": "LDG0747 m3", "ghi_chu": "t3", "index": "113"},
    {"file": "LDG0747_m3_t1a", "template": "kieu_B", "ngay": "15/07/2026",
     "lat": "11.69111174", "lon": "108.37475665", "bearing_deg": "28", "huong": "28° NE",
     "huyen": "Đức Trọng", "tinh": "Lâm Đồng", "altitude": "903.8", "speed": "0.0",
     "ma_tram": "LDG0747 m3", "ghi_chu": "t1", "index": "135"},
    {"file": "LDG0747_m3_t1b", "template": "kieu_B", "ngay": "15/07/2026",
     "lat": "11.69118505", "lon": "108.37464451", "bearing_deg": "131", "huong": "131° SE",
     "huyen": "Đức Trọng", "tinh": "Lâm Đồng", "altitude": "903.5", "speed": "0.0",
     "ma_tram": "LDG0747 m3", "ghi_chu": "t1", "index": "117"},
]


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "du_lieu"
    fill = PatternFill("solid", fgColor="1F4E78")
    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = max(10, len(h) + 4)
    for r, row in enumerate(SAMPLES, 2):
        for c, h in enumerate(HEADERS, 1):
            ws.cell(r, c, row.get(h, ""))
    ws.freeze_panes = "A2"
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print(f"OK -> {OUT}")


if __name__ == "__main__":
    main()
